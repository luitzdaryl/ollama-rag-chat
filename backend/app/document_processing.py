import csv
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
import re

def clean_text(text: str) -> str:
    """Strip common PDF extraction artifacts — TOC dot leaders, horizontal
    rules, and similar repeated-character noise that carries no real meaning."""
    text = re.sub(r'[.\-_]{4,}', ' ', text)   # dot leaders, dashed rules, underscores
    text = re.sub(r'\n{3,}', '\n\n', text)     # collapse excessive blank lines
    return text

def extract_text(file_path: str) -> str:
    return clean_text(_extract_text_raw(file_path))

def _extract_text_raw(file_path: str) -> str:
    """Extract plain text from a file, dispatching by extension."""
    suffix = Path(file_path).suffix.lower()

    if suffix in (".txt", ".md"):
        return Path(file_path).read_text(encoding="utf-8")

    if suffix == ".pdf":
        reader = PdfReader(file_path)
        pages_text = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages_text)

    if suffix == ".docx":
        doc = Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    if suffix == ".xlsx":
        wb = load_workbook(file_path, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)

    if suffix == ".csv":
        lines = []
        with open(file_path, newline="", encoding="utf-8") as f:
            for row in csv.reader(f):
                if row:
                    lines.append(" | ".join(row))
        return "\n".join(lines)

    raise ValueError(f"Unsupported file type: {suffix}")